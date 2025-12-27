import os
import streamlit as st
from dotenv import load_dotenv
from typing import List
from difflib import SequenceMatcher

from langchain_openai import ChatOpenAI, OpenAIEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.output_parsers import StrOutputParser
from langchain_core.documents import Document

from pypdf import PdfReader
import docx
import numpy as np

# --------------------------------------------------
# ENV
# --------------------------------------------------
load_dotenv()

# --------------------------------------------------
# STREAMLIT UI
# --------------------------------------------------
st.set_page_config(page_title="AI Document Chat", layout="centered")
st.title("📄 AI Document Chat")
st.caption("Upload PDFs or DOCX files and ask questions")

# --------------------------------------------------
# SESSION STATE
# --------------------------------------------------
if "vectorstore" not in st.session_state:
    st.session_state.vectorstore = None

if "chat_history" not in st.session_state:
    st.session_state.chat_history = []

if "qa_cache" not in st.session_state:
    st.session_state.qa_cache = {}

# --------------------------------------------------
# UTILITIES
# --------------------------------------------------
def normalize_question(q: str) -> str:
    return q.lower().strip()

def find_similar_cached_question(q: str, threshold: float = 0.9):
    for cached_q in st.session_state.qa_cache:
        if SequenceMatcher(None, q, cached_q).ratio() >= threshold:
            return cached_q
    return None

# --------------------------------------------------
# FILE TEXT EXTRACTION
# --------------------------------------------------
def extract_text_from_pdf(file) -> str:
    reader = PdfReader(file)
    return "\n".join(page.extract_text() or "" for page in reader.pages)

def extract_text_from_docx(file) -> str:
    doc = docx.Document(file)
    return "\n".join(p.text for p in doc.paragraphs)

# --------------------------------------------------
# VECTOR STORE
# --------------------------------------------------
def build_vectorstore(texts: List[str]):
    documents = [Document(page_content=t) for t in texts]

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=150
    )
    chunks = splitter.split_documents(documents)

    embeddings = OpenAIEmbeddings()
    return FAISS.from_documents(chunks, embeddings)

# --------------------------------------------------
# FILE UPLOAD
# --------------------------------------------------
uploaded_files = st.file_uploader(
    "Upload PDF or DOCX files",
    type=["pdf", "docx"],
    accept_multiple_files=True
)

if uploaded_files and st.session_state.vectorstore is None:
    with st.spinner("Processing documents..."):
        texts = []
        for file in uploaded_files:
            if file.name.endswith(".pdf"):
                texts.append(extract_text_from_pdf(file))
            elif file.name.endswith(".docx"):
                texts.append(extract_text_from_docx(file))

        st.session_state.vectorstore = build_vectorstore(texts)

    st.success("Knowledge base created successfully")

# --------------------------------------------------
# LLM + PROMPT
# --------------------------------------------------
llm = ChatOpenAI(model="gpt-4o-mini", temperature=0)

prompt = ChatPromptTemplate.from_template("""
You are a document-based QA system.

Rules:
- Answer ONLY using the provided context.
- Do NOT use outside knowledge.
- If information is incomplete or missing, say so clearly.

Context:
{context}

Question:
{question}
""")

chain = prompt | llm | StrOutputParser()

# --------------------------------------------------
# CHAT UI
# --------------------------------------------------
st.divider()
st.subheader("Chat")

for role, message in st.session_state.chat_history:
    st.chat_message(role).write(message)

user_query = st.chat_input("Ask a question about the uploaded documents")

if user_query:
    if st.session_state.vectorstore is None:
        st.error("Please upload documents first.")
    else:
        normalized_q = normalize_question(user_query)

        cached_match = (
            normalized_q
            if normalized_q in st.session_state.qa_cache
            else find_similar_cached_question(normalized_q)
        )

        st.session_state.chat_history.append(("user", user_query))

        if cached_match:
            st.session_state.chat_history.append(
                ("assistant", st.session_state.qa_cache[cached_match])
            )
            st.rerun()

        # -------- RETRIEVAL WITH CONFIDENCE --------
        retriever = st.session_state.vectorstore.similarity_search_with_score(
            user_query, k=6
        )

        docs, scores = zip(*retriever)
        avg_distance = float(np.mean(scores))

        HIGH_CONF = 1.6
        MID_CONF = 2.0

        if avg_distance > MID_CONF:
            answer = (
                "I cannot find this information clearly in the uploaded documents. "
                "Please rephrase or ask something more specific."
            )
        else:
            context = "\n\n".join(doc.page_content for doc in docs)[:3500]

            answer = chain.invoke({
                "context": context,
                "question": user_query
            })

            # cache ONLY confident answers
            if avg_distance <= HIGH_CONF:
                st.session_state.qa_cache[normalized_q] = answer

        st.session_state.chat_history.append(("assistant", answer))
        st.rerun()